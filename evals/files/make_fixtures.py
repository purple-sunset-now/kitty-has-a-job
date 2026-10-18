"""Builds the test CVs used by the evals. Run: python3 evals/files/make_fixtures.py"""
import os, shutil, subprocess, tempfile, zipfile
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))

def bullet(doc, lead, rest):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(lead + " "); r.bold = True
    p.add_run(rest)

def base_cv(extra_roles=0):
    d = Document()
    d.add_heading("Alex Example", 0)
    d.add_paragraph("Berlin, Germany | alex@example.com")
    d.add_heading("Summary", 1)
    d.add_paragraph("Senior product manager with 9 years in B2B SaaS and marketplaces. "
                    "Shipped ML-powered search and pricing products used by 3M users. "
                    "Works hands-on with data and engineering teams.")
    d.add_heading("Experience", 1)
    d.add_paragraph("Senior Product Manager, Shoply (2022 to now)")
    bullet(d, "Launched", "ML search ranking that lifted conversion by 9% for 3M monthly users.")
    bullet(d, "Cut", "checkout drop-off by 14% by redesigning payment flows with 2 teams.")
    bullet(d, "Built", "pricing experiments platform used by 6 teams, 40 tests per quarter.")
    d.add_paragraph("Product Manager, Rentco (2018 to 2022)")
    bullet(d, "Grew", "B2B marketplace supply from 200 to 1,400 partners in two years.")
    bullet(d, "Responsible", "for various partner tools and helped with onboarding.")
    for i in range(extra_roles):
        d.add_paragraph(f"Product Manager, Filler Co {i} (2010 to 2018)")
        for j in range(6):
            bullet(d, "Delivered", f"project {i}-{j} with a long description that fills space on the page " * 2)
    d.add_heading("Education", 1)
    d.add_paragraph("BSc Computer Science")
    return d

def add_hidden_bullet_and_customxml(path, text):
    tmp = tempfile.mkdtemp()
    with zipfile.ZipFile(path) as z: z.extractall(tmp)
    p = os.path.join(tmp, "word", "document.xml")
    s = open(p, encoding="utf-8").read()
    sdt = ('<w:sdt><w:sdtContent><w:p><w:pPr><w:pStyle w:val="ListBullet"/></w:pPr>'
           f'<w:r><w:rPr><w:b/></w:rPr><w:t xml:space="preserve">Led </w:t></w:r><w:r><w:t>{text}</w:t></w:r></w:p>'
           '</w:sdtContent></w:sdt>')
    anchor = "Product Manager, Rentco"
    i = s.index(anchor); i = s.rfind("<w:p>", 0, i) if s.rfind("<w:p>", 0, i) != -1 else s.rfind("<w:p ", 0, i)
    s = s[:i] + sdt + s[i:]
    open(p, "w", encoding="utf-8").write(s)
    os.makedirs(os.path.join(tmp, "customXml"), exist_ok=True)
    open(os.path.join(tmp, "customXml", "item1.xml"), "w").write('<?xml version="1.0"?><root/>')
    os.remove(path)
    subprocess.run(["zip", "-q", "-r", "-X", path, "."], cwd=tmp, check=True)
    shutil.rmtree(tmp)

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def patch(path, fn, extra_files=None):
    """Unzip, let fn edit document.xml (and other parts), rezip."""
    tmp = tempfile.mkdtemp()
    with zipfile.ZipFile(path) as z: z.extractall(tmp)
    fn(tmp)
    for name, data in (extra_files or {}).items():
        os.makedirs(os.path.dirname(os.path.join(tmp, name)), exist_ok=True)
        open(os.path.join(tmp, name), "w", encoding="utf-8").write(data)
    os.remove(path)
    subprocess.run(["zip", "-q", "-r", "-X", path, "."], cwd=tmp, check=True)
    shutil.rmtree(tmp)


def edit_xml(tmp, part, fn):
    p = os.path.join(tmp, part)
    s = open(p, encoding="utf-8").read()
    open(p, "w", encoding="utf-8").write(fn(s))


def glyph_cv(path):
    """Bullets typed as a '•' character and a tab, no Word numbering."""
    d = Document()
    d.add_heading("Alex Example", 0)
    d.add_paragraph("Berlin, Germany | alex@example.com")
    d.add_heading("Summary", 1)
    d.add_paragraph("Senior product manager with 9 years in B2B SaaS and marketplaces. "
                    "Shipped ML-powered search and pricing products used by 3M users.")
    d.add_heading("Experience", 1)
    def gb(lead, rest):
        p = d.add_paragraph()
        p.add_run("\u2022\t")
        r = p.add_run(lead + " "); r.bold = True
        p.add_run(rest)
    d.add_paragraph("Senior Product Manager, Shoply (2022 to now)")
    gb("Launched", "ML search ranking that lifted conversion by 9% for 3M monthly users.")
    gb("Cut", "checkout drop-off by 14% by redesigning payment flows with 2 teams.")
    gb("Built", "pricing experiments platform used by 6 teams, 40 tests per quarter.")
    d.add_paragraph("Product Manager, Rentco (2018 to 2022)")
    gb("Grew", "B2B marketplace supply from 200 to 1,400 partners in two years.")
    gb("Responsible", "for various partner tools and helped with onboarding.")
    d.add_heading("Education", 1)
    d.add_paragraph("BSc Computer Science")
    d.save(path)


def table_cv(path):
    """Experience laid out in a two-column table: dates left, role and bullets right."""
    d = Document()
    d.add_heading("Alex Example", 0)
    d.add_paragraph("Berlin, Germany | alex@example.com")
    d.add_heading("Summary", 1)
    d.add_paragraph("Senior product manager with 9 years in B2B SaaS and marketplaces. "
                    "Shipped ML-powered search and pricing products used by 3M users.")
    d.add_heading("Experience", 1)
    t = d.add_table(rows=2, cols=2)
    roles = [("2022 to now", "Senior Product Manager, Shoply",
              [("Launched", "ML search ranking that lifted conversion by 9% for 3M monthly users."),
               ("Cut", "checkout drop-off by 14% by redesigning payment flows with 2 teams."),
               ("Built", "pricing experiments platform used by 6 teams, 40 tests per quarter.")]),
             ("2018 to 2022", "Product Manager, Rentco",
              [("Grew", "B2B marketplace supply from 200 to 1,400 partners in two years."),
               ("Responsible", "for various partner tools and helped with onboarding.")])]
    for row, (dates, title, bullets) in zip(t.rows, roles):
        row.cells[0].text = dates
        c = row.cells[1]
        c.paragraphs[0].add_run(title).bold = True
        for lead, rest in bullets:
            p = c.add_paragraph(style="List Bullet")
            p.add_run(lead + " ").bold = True
            p.add_run(rest)
    d.add_heading("Education", 1)
    d.add_paragraph("BSc Computer Science")
    d.save(path)


def german_cv(path):
    """German CV; Word has split sentences across runs (proofing marks), as real files often do."""
    d = Document()
    d.add_heading("Alex Example", 0)
    d.add_paragraph("Berlin | alex@example.com")
    d.add_heading("Profil", 1)
    d.add_paragraph("SPLIT_SUMMARY")
    d.add_heading("Berufserfahrung", 1)
    d.add_paragraph("Senior Product Manager, Shoply (seit 2022)")
    bullet(d, "Eingeführt:", "ML-Suchranking, das die Conversion bei 3 Mio. Nutzern pro Monat um 9 % steigerte.")
    bullet(d, "Gesenkt:", "Checkout-Abbrüche um 14 % durch neue Zahlungsabläufe mit 2 Teams.")
    bullet(d, "Aufgebaut:", "Plattform für Preisexperimente, genutzt von 6 Teams, 40 Tests pro Quartal.")
    d.add_paragraph("Product Manager, Rentco (2018 bis 2022)")
    bullet(d, "Gesteigert:", "Anbieter im B2B-Marktplatz von 200 auf 1.400 Partner in zwei Jahren.")
    bullet(d, "Verantwortlich", "SPLIT_BULLET")
    d.add_heading("Ausbildung", 1)
    d.add_paragraph("BSc Informatik")
    d.save(path)
    def split(tmp):
        def fn(s):
            summ = ('Senior Product Manager mit 9 Jahren Erfahrung in B2B-SaaS und </w:t></w:r>'
                    '<w:proofErr w:type="spellStart"/><w:r w:rsidR="00A1B2C3"><w:t>Marktplätzen</w:t></w:r>'
                    '<w:proofErr w:type="spellEnd"/><w:r><w:t xml:space="preserve">. Hat ML-gestützte Such- und Preisprodukte für 3 Mio. Nutzer gelauncht.')
            s = s.replace("SPLIT_SUMMARY", summ)
            bul = ('für verschiedene </w:t></w:r><w:proofErr w:type="spellStart"/>'
                   '<w:r w:rsidR="00D4E5F6"><w:t>Partner-Tools</w:t></w:r><w:proofErr w:type="spellEnd"/>'
                   '<w:r><w:t xml:space="preserve"> und half beim Onboarding.')
            return s.replace("SPLIT_BULLET", bul)
        edit_xml(tmp, "word/document.xml", fn)
    patch(path, split)


COMMENTS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:comments {W}><w:comment w:id="0" w:author="Reviewer" w:date="2026-09-01T10:00:00Z" w:initials="R">'
            '<w:p><w:r><w:t>Do we have a number for this? (private reviewer note)</w:t></w:r></w:p></w:comment></w:comments>')


def tracked_cv(path):
    """Template with a tracked change, a reviewer comment and the contact line in a text box."""
    base_cv().save(path)
    def fn(tmp):
        def doc(s):
            s = s.replace("Works hands-on with data and engineering teams.",
                          'Works hands-on with data and engineering teams.</w:t></w:r>'
                          '<w:del w:id="91" w:author="Reviewer" w:date="2026-09-01T10:00:00Z"><w:r><w:delText xml:space="preserve"> Loves spreadsheets.</w:delText></w:r></w:del>'
                          '<w:ins w:id="92" w:author="Reviewer" w:date="2026-09-01T10:00:00Z"><w:r><w:t xml:space="preserve"> Based in Berlin.</w:t></w:r></w:ins>'
                          '<w:r><w:t xml:space="preserve">')
            g = "B2B marketplace supply from 200 to 1,400 partners in two years."
            i = s.index(g); ps = s.rfind("<w:p>", 0, i); pe = s.index("</w:p>", i)
            para = s[ps:pe]
            para = para.replace("<w:r>", '<w:commentRangeStart w:id="0"/><w:r>', 1) + \
                '<w:commentRangeEnd w:id="0"/><w:r><w:commentReference w:id="0"/></w:r>'
            s = s[:ps] + para + s[pe:]
            box = ('<w:p><w:r><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml" style="width:300pt;height:24pt">'
                   '<v:textbox><w:txbxContent><w:p><w:r><w:t>Phone +49 30 0000000 | linkedin.com/in/alex-example</w:t></w:r></w:p>'
                   '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>')
            j = s.index("Berlin, Germany | alex@example.com"); j = s.index("</w:p>", j) + len("</w:p>")
            return s[:j] + box + s[j:]
        edit_xml(tmp, "word/document.xml", doc)
        edit_xml(tmp, "word/_rels/document.xml.rels", lambda s: s.replace("</Relationships>",
            '<Relationship Id="rIdC1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="comments.xml"/></Relationships>'))
        edit_xml(tmp, "[Content_Types].xml", lambda s: s.replace("</Types>",
            '<Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/></Types>'))
    patch(path, fn, {"word/comments.xml": COMMENTS})


os.makedirs(os.path.join(HERE, "templates"), exist_ok=True)
a = os.path.join(HERE, "templates", "cv_main.docx")
base_cv().save(a)
add_hidden_bullet_and_customxml(a, "discovery for an AI assistant for sellers, adopted by 12,000 sellers in 6 months.")
base_cv(extra_roles=6).save(os.path.join(HERE, "cv_too_long.docx"))
glyph_cv(os.path.join(HERE, "cv_glyph_bullets.docx"))
table_cv(os.path.join(HERE, "cv_table.docx"))
german_cv(os.path.join(HERE, "cv_german.docx"))
tracked_cv(os.path.join(HERE, "cv_tracked.docx"))
print("fixtures written")
